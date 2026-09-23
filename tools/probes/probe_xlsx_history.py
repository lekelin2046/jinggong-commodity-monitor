#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""探查《诺博橡胶大宗物料价格走势-2026.xlsx》历史序列 → 与 data.json 比对
只读探查，不写任何文件。
"""
import json
import datetime as dt
from pathlib import Path

from openpyxl import load_workbook

# 2026-09-23 整理：源文件已归入 sources/橡胶/，路径锚定仓库根（不再依赖 cwd）
_ROOT = Path(__file__).resolve().parents[2]
XLSX = str(_ROOT / "sources" / "橡胶" / "诺博橡胶大宗物料价格走势-2026.xlsx")
DATA = str(_ROOT / "docs" / "plastic" / "data.json")

# (sheet, col_letter) -> code
MAP = {
    ("铝 -日", "B"): "AL_A00",
    ("铝 -日", "C"): "AL_A380",
    ("铝 -日", "D"): "AL_ZLD104",
    ("铝 -日", "E"): "AL_ALSI9CU3",
    ("三元乙丙-月", "B"): "C2_ETHYLENE",
    ("三元乙丙-月", "C"): "C3_PROPYLENE",
    ("三元乙丙-月", "D"): "EPDM_6950C",
    ("三元乙丙-月", "E"): "EPDM_3110M",
    ("三元乙丙-月", "F"): "EPDM_13561C",
    ("炭黑-月", "B"): "CB_N550",
    ("炭黑-月", "C"): "CB_COALTAR",
    ("天然胶-月", "B"): "NR_SCRWF",
    ("天然胶-月", "C"): "NR_RSS3",
    ("天然胶-月", "E"): "BR9000",
    ("促进剂防老剂", "B"): "ACC_DM",
    ("促进剂防老剂", "C"): "ACC_CZ",
    ("促进剂防老剂", "D"): "ACC_M",
    ("促进剂防老剂", "E"): "ACC_TMTD",
    ("促进剂防老剂", "F"): "AO_4020",
}


def parse_date(v):
    """返回 datetime.date 或 None"""
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        # Excel 序列号（1900 日期系统）
        try:
            n = float(v)
            if 20000 < n < 60000:
                return (dt.date(1899, 12, 30) + dt.timedelta(days=int(n)))
        except Exception:
            return None
        return None
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return f if f > 0 else None
    s = str(v).strip().replace(",", "")
    if not s or not any(ch.isdigit() for ch in s):
        return None
    try:
        f = float(s)
        return f if f > 0 else None
    except ValueError:
        return None


wb = load_workbook(XLSX, data_only=True, read_only=True)
board = json.load(open(DATA, encoding="utf-8"))["data"]
codes = ["AL_A00", "AL_A380", "AL_ZLD104", "AL_ALSI9CU3", "C2_ETHYLENE", "C3_PROPYLENE",
         "EPDM_6950C", "EPDM_3110M", "EPDM_13561C", "CB_N550", "CB_COALTAR",
         "NR_SCRWF", "NR_RSS3", "BR9000", "ACC_DM", "ACC_CZ", "ACC_M", "ACC_TMTD", "AO_4020"]
board_cov = {c: sum(1 for v in board.values() if isinstance(v, dict) and v.get(c) not in (None, ""))
             for c in codes}

series = {}
for (sh, col), code in MAP.items():
    ws = wb[sh]
    d = {}
    dup = 0
    for row in ws.iter_rows(min_col=1, max_col=ws.max_column):
        a = row[0].value if len(row) > 0 else None
        dte = parse_date(a)
        if dte is None:
            continue
        try:
            cidx = ord(col) - 65
        except Exception:
            continue
        if cidx >= len(row):
            continue
        v = num(row[cidx].value)
        if v is None:
            continue
        key = dte.isoformat()
        if key in d:
            dup += 1
        d[key] = v
    series[code] = (d, dup)

print("=" * 104)
print("Excel 历史序列探查 vs data.json 现状")
print("=" * 104)
print(f"{'code':<14}{'Excel点':>8}{'板内点':>8}{'重叠':>7}{'板缺可补':>9}{'新增日期行':>10}{'周末点':>7}{'重复日':>7}{'首次':>12}{'末次':>12}")
print("-" * 104)
all_new_dates = set()
summary = {}
for c in codes:
    d, dup = series[c]
    ov = [k for k in d if k in board]
    miss_fill = [k for k in d if k in board and board[k].get(c) in (None, "")]
    new_rows = [k for k in d if k not in board]
    wk = [k for k in d if dt.date.fromisoformat(k).weekday() >= 5]
    all_new_dates |= set(new_rows)
    summary[c] = dict(excel=len(d), board=board_cov[c], overlap=len(ov),
                      fill=len(miss_fill), new_rows=len(new_rows), weekend=len(wk), dup=dup,
                      first=min(d) if d else "-", last=max(d) if d else "-")
    print(f"{c:<14}{len(d):>8}{board_cov[c]:>8}{len(ov):>7}{len(miss_fill):>9}{len(new_rows):>10}"
          f"{len(wk):>7}{dup:>7}{min(d) if d else '-':>12}{max(d) if d else '-':>12}")

print()
print("合并后预估 total_days =", len(board) + len(all_new_dates), f"(现 {len(board)}, 新增日期行 {len(all_new_dates)})")
print("最早新增日期 =", min(all_new_dates) if all_new_dates else "-")

# 重叠区间的口径差异（判断是否会造成接缝跳变）
print()
print("=" * 104)
print("重叠日期上的数值差异（板内已抓取值 vs Excel 值）—— insert-only 会保留板内值，\"接缝\"风险在此")
print("=" * 104)
print(f"{'code':<14}{'重叠':>6}{'完全一致':>9}{'差异':>7}{'最大绝对差':>12}{'最大相对差':>11}{'样例(日期:板/Excel)':>34}")
print("-" * 104)
for c in codes:
    d, _ = series[c]
    diffs = []
    for k in d:
        if k in board:
            bv = board[k].get(c)
            if bv not in (None, ""):
                try:
                    bf = float(bv)
                except (TypeError, ValueError):
                    continue
                if abs(bf - d[k]) > 1e-6:
                    diffs.append((abs(bf - d[k]), abs(bf - d[k]) / (abs(d[k]) or 1), k, bf, d[k]))
    ov = [k for k in d if k in board and board[k].get(c) not in (None, "")]
    if not ov:
        print(f"{c:<14}{0:>6}{'-':>9}{'-':>7}{'-':>12}{'-':>11}{'-':>34}")
        continue
    if diffs:
        diffs.sort(reverse=True)
        mx = diffs[0]
        sample = f"{mx[2]}:{mx[3]:g}/{mx[4]:g}"
        print(f"{c:<14}{len(ov):>6}{len(ov)-len(diffs):>9}{len(diffs):>7}{mx[0]:>12.2f}{mx[1]*100:>10.2f}%{sample:>34}")
    else:
        print(f"{c:<14}{len(ov):>6}{len(ov):>9}{0:>7}{0:>12.2f}{0:>10.2f}%{'全部一致':>34}")
