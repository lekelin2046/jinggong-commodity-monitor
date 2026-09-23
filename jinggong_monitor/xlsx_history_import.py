#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""《诺博橡胶大宗物料价格走势-2026.xlsx》历史数据 → docs/plastic/data.json

用途
    甲方提供的年度物料价格走势工作簿里含多年历史（2020 起），而看板部分品种
    （铝 4 项、高温煤焦油、促进剂 DM/CZ、顺丁橡胶）只有最近几个月的抓取值。
    本模块把工作簿中的历史序列按品种映射后**增量**并入 data.json。

默认策略 gap（只补空缺，绝不覆盖）
    仅当 data.json 该日期该品种为空时才写入。已在库里的抓取值（SMM/隆众/百川实时口径）
    一律保留 —— 抓取值是已对外发布的运行口径，不被文档覆盖。
    --excel-priority 可切换为「工作簿优先」（覆盖重叠区间），会改变已发布值，慎用。

三重守卫
    1. 日历守卫：只写「交易日」（含调休补班日）或 data.json 中已存在的日期，
       不因工作簿里的周末顺延行而新增周末列。
    2. 离群守卫：与同序列相邻 ±10 点中位数偏离 >50% 的点判为录入错误，跳过并报告
       （实测工作簿存在 2021-12-06 EPDM6950C=229000 这类孤点）。
    3. 不替代守卫：品种在工作簿中找不到对应列时不猜测、不用近似品种顶替
       （SBR1712 丁苯橡胶 1712 工作簿只有 1502 牌号 → 不导入）。

用法
    python -m jinggong_monitor.xlsx_history_import --report     # 只出体检报告，不写盘
    python -m jinggong_monitor.xlsx_history_import --apply      # 体检通过后写盘（自动备份）
    python -m jinggong_monitor.xlsx_history_import --apply --excel-priority
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import shutil
import statistics
import sys

from openpyxl import load_workbook

from .trading_calendar import is_trading_day

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# 2026-09-23 整理：源工作簿已归入 sources/橡胶/
XLSX_DEFAULT = os.path.join(BASE, "sources", "橡胶", "诺博橡胶大宗物料价格走势-2026.xlsx")
DATA_DEFAULT = os.path.join(BASE, "docs", "plastic", "data.json")

# (工作表, 列字母) -> 看板品种代码
MAPPING: dict[tuple[str, str], str] = {
    ("铝 -日", "B"): "AL_A00",
    ("铝 -日", "C"): "AL_A380",
    ("铝 -日", "D"): "AL_ZLD104",
    ("铝 -日", "E"): "AL_ALSI9CU3",
    ("三元乙丙-月", "B"): "C2_ETHYLENE",
    ("三元乙丙-月", "C"): "C3_PROPYLENE",
    ("三元乙丙-月", "D"): "EPDM_6950C",
    ("三元乙丙-月", "E"): "EPDM_3110M",
    ("三元乙丙-月", "F"): "EPDM_13561C",
    ("炭黑-月", "B"): "CB_N550",
    ("炭黑-月", "C"): "CB_COALTAR",
    ("天然胶-月", "B"): "NR_SCRWF",
    ("天然胶-月", "C"): "NR_RSS3",
    ("天然胶-月", "E"): "BR9000",
    ("促进剂防老剂", "B"): "ACC_DM",
    ("促进剂防老剂", "C"): "ACC_CZ",
    ("促进剂防老剂", "D"): "ACC_M",
    ("促进剂防老剂", "E"): "ACC_TMTD",
    ("促进剂防老剂", "F"): "AO_4020",
}

# 工作簿中无对应列、明确不导入的监控品种（不猜、不顶替）
UNMATCHED = {
    "SBR1712": "工作簿无 1712 牌号（仅「天然胶-月」D 列丁苯胶1502，牌号不符，不顶替）",
    "OIL_WTI": "工作簿为布伦特原油，看板 OIL_WTI 为 WTI（akshare），基准不同，不混接",
}

OUTLIER_RATIO = 0.50   # 与相邻中位数偏离超过此比例 → 判为录入错误
NEIGHBOURS = 10        # 取相邻点数（前后各 10）


# ---------------------------------------------------------------- 解析

def parse_date(v):
    """工作簿日期混用「真日期」与「文本」，统一解析为 date；失败返回 None。"""
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        n = float(v)
        if 20000 < n < 60000:          # Excel 1900 日期序列号
            return dt.date(1899, 12, 30) + dt.timedelta(days=int(n))
        return None
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d",
                "%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def num(v):
    """取正数值；非数值/空/汇总文本返回 None。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return f if f > 0 else None
    s = str(v).strip().replace(",", "")
    if not s or not any(ch.isdigit() for ch in s):
        return None
    # 过滤「周均价」「月均价」「年均价」等汇总行残留
    if any(t in s for t in ("均价", "平均", "同比", "环比", "%")):
        return None
    try:
        f = float(s)
        return f if f > 0 else None
    except ValueError:
        return None


def extract(xlsx: str = XLSX_DEFAULT):
    """返回 ({code: {iso_date: value}}, {code: 解析统计})。"""
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    series: dict[str, dict[str, float]] = {}
    stats: dict[str, dict] = {}
    for (sheet, col), code in MAPPING.items():
        if sheet not in wb.sheetnames:
            raise KeyError(f"工作表缺失：{sheet}")
        ws = wb[sheet]
        ci = ord(col) - ord("A")
        d: dict[str, float] = {}
        dup_conflict = 0
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                               max_col=ci + 1, values_only=True):
            dte = parse_date(row[0] if len(row) > 0 else None)
            if dte is None:
                continue
            v = num(row[ci]) if ci < len(row) else None
            if v is None:
                continue
            key = dte.isoformat()
            if key in d and abs(d[key] - v) > 1e-6:
                dup_conflict += 1
                d[key] = v          # 同日重复：以表中靠后的行为准
            else:
                d[key] = v
        series[code] = d
        stats[code] = {"sheet": sheet, "col": col, "points": len(d),
                       "first": min(d) if d else None, "last": max(d) if d else None,
                       "dup_conflict": dup_conflict}
    return series, stats


# ---------------------------------------------------------------- 校验

def find_outliers(d: dict[str, float]) -> list[tuple[str, float, float]]:
    """孤立离群点检测：与相邻 ±NEIGHBOURS 点中位数偏离 >OUTLIER_RATIO 者。

    返回 [(日期, 值, 邻域中位数), ...]。
    """
    keys = sorted(d)
    out = []
    for i, k in enumerate(keys):
        lo = max(0, i - NEIGHBOURS)
        hi = min(len(keys), i + NEIGHBOURS + 1)
        nb = [d[keys[j]] for j in range(lo, hi) if j != i]
        if len(nb) < 5:
            continue
        med = statistics.median(nb)
        if med <= 0:
            continue
        if abs(d[k] - med) / med > OUTLIER_RATIO:
            out.append((k, d[k], med))
    return out


def build_plan(series, board, policy="gap"):
    """生成写入计划。返回 (plan, report)。

    plan: {iso_date: {code: value}}
    """
    plan: dict[str, dict[str, float]] = {}
    report: dict[str, dict] = {}

    for code, d in series.items():
        outliers = find_outliers(d)
        out_keys = {k for k, _, _ in outliers}

        fill, overwrite, skipped_offcal, skipped_outlier = [], [], [], []
        for k, v in d.items():
            if k in out_keys:
                skipped_outlier.append(k)
                continue
            row = board.get(k)
            has_board_val = bool(row) and row.get(code) not in (None, "")
            if has_board_val:
                if policy == "excel":
                    # 仅覆盖工作簿覆盖范围内、且确实不同的点
                    try:
                        same = abs(float(row[code]) - v) <= 1e-6
                    except (TypeError, ValueError):
                        same = False
                    if not same:
                        overwrite.append(k)
                        plan.setdefault(k, {})[code] = v
                continue
            # 该品种在此日期无值 → 补
            if k not in board:
                dte = dt.date.fromisoformat(k)
                if not is_trading_day(dte):
                    skipped_offcal.append(k)
                    continue
            fill.append(k)
            plan.setdefault(k, {})[code] = v

        report[code] = {
            "excel_points": len(d),
            "board_points": sum(1 for r in board.values()
                                if isinstance(r, dict) and r.get(code) not in (None, "")),
            "fill": len(fill), "overwrite": len(overwrite),
            "skip_offcal": len(skipped_offcal), "skip_outlier": len(skipped_outlier),
            "outliers": outliers,
            "fill_first": min(fill) if fill else None,
            "fill_last": max(fill) if fill else None,
        }
    return plan, report


# ---------------------------------------------------------------- 输出

def print_report(series, stats, board, plan, report, policy):
    print("=" * 108)
    print(f"工作簿历史导入体检  策略={policy}   "
          f"{'（只补空缺，不覆盖已抓取值）' if policy == 'gap' else '（工作簿优先，会覆盖重叠区间）'}")
    print("=" * 108)
    print(f"{'code':<14}{'工作簿点':>9}{'库内点':>8}{'可补缺口':>9}{'覆盖':>7}"
          f"{'跳非交易日':>11}{'跳离群':>8}{'补首':>13}{'补末':>13}")
    print("-" * 108)
    for code in sorted(report):
        r = report[code]
        print(f"{code:<14}{r['excel_points']:>9}{r['board_points']:>8}{r['fill']:>9}"
              f"{r['overwrite']:>7}{r['skip_offcal']:>11}{r['skip_outlier']:>8}"
              f"{str(r['fill_first'] or '-'):>13}{str(r['fill_last'] or '-'):>13}")

    new_dates = [k for k in plan if k not in board]
    print("-" * 108)
    print(f"合计写入数据点 = {sum(len(v) for v in plan.values())}"
          f" | 新增日期行 = {len(new_dates)}"
          f" | 预估 total_days = {len(board) + len(new_dates)}（现 {len(board)}）")
    if new_dates:
        print(f"新增日期区间 = {min(new_dates)} ~ {max(new_dates)}")

    all_out = [(c, k, v, m) for c, r in report.items() for (k, v, m) in r["outliers"]]
    if all_out:
        print()
        print("⚠️ 离群点（判为工作簿录入错误，已跳过，不入库）：")
        for c, k, v, m in all_out:
            print(f"   {c:<14}{k}  值={v:>12,.2f}  邻域中位数={m:>12,.2f}  "
                  f"偏离={(abs(v - m) / m * 100):.1f}%")

    print()
    print("未导入的监控品种（工作簿中无对应列，未做任何替代）：")
    for code, why in UNMATCHED.items():
        print(f"   {code:<14}{why}")

    # 接缝提示：补入历史后，与库内既有值的首日差异
    seams = []
    for code, d in series.items():
        r = report[code]
        if not r["fill_first"]:
            continue
        later = sorted(k for k in d if k in board
                       and isinstance(board[k], dict)
                       and board[k].get(code) not in (None, ""))
        if not later:
            continue
        k0 = later[0]
        try:
            bv = float(board[k0][code])
        except (TypeError, ValueError):
            continue
        ev = d.get(k0)
        if ev and abs(bv - ev) / abs(ev) > 0.02:
            seams.append((code, r["fill_last"], k0, bv, ev, (bv - ev) / ev * 100))
    if seams:
        print()
        print("⚠️ 口径接缝（历史段来自工作簿、该日起为库内抓取值，两段存在系统性差异）：")
        for code, fl, k0, bv, ev, pct in seams:
            print(f"   {code:<14}工作簿段止于 {fl} → 库内值起于 {k0}："
                  f"库内 {bv:,.0f} vs 工作簿 {ev:,.0f}（{pct:+.2f}%）")


def main(argv=None):
    ap = argparse.ArgumentParser(description="工作簿历史数据导入 data.json")
    ap.add_argument("--report", action="store_true", help="只体检不写盘（默认）")
    ap.add_argument("--apply", action="store_true", help="写盘（自动备份）")
    ap.add_argument("--excel-priority", action="store_true", help="工作簿优先（覆盖重叠区间）")
    ap.add_argument("--xlsx", default=XLSX_DEFAULT)
    ap.add_argument("--data", default=DATA_DEFAULT)
    args = ap.parse_args(argv)

    policy = "excel" if args.excel_priority else "gap"
    series, stats = extract(args.xlsx)
    with open(args.data, encoding="utf-8") as f:
        doc = json.load(f)
    board = doc["data"]

    plan, report = build_plan(series, board, policy)
    print_report(series, stats, board, plan, report, policy)

    if not args.apply:
        print()
        print("（体检模式，未写盘。确认无误后加 --apply 执行）")
        return 0

    if not plan:
        print()
        print("无任何可写入的数据点 → 跳过写盘。")
        return 0

    # 2026-09-23 整理：备份统一落 backups/塑料/，不再堆在 docs/plastic/ 下
    bak_dir = os.path.join(BASE, "backups", "塑料")
    os.makedirs(bak_dir, exist_ok=True)
    bak = os.path.join(bak_dir, os.path.basename(args.data) + ".bak-"
                       + dt.datetime.now().strftime("%Y%m%d-%H%M%S"))
    shutil.copy2(args.data, bak)
    print()
    print(f"已备份 → {os.path.relpath(bak, BASE)}")

    added = 0
    for k, vals in plan.items():
        row = board.setdefault(k, {})
        for code, v in vals.items():
            if row.get(code) in (None, ""):
                added += 1
            row[code] = v
    doc["data"] = board
    doc["total_days"] = len(board)
    doc["last_updated"] = max(board)
    with open(args.data, "w", encoding="utf-8") as f:
        # 必须与 export_excel_to_json.py 一致：1 空格缩进、无尾换行，
        # 否则整文件重排、diff 失真（已实测往返字节一致）
        json.dump(doc, f, ensure_ascii=False, indent=1)
    print(f"✓ 写入完成：新增 {added} 个数据点 | 覆盖 {sum(len(v) for v in plan.values()) - added} 个"
          f" | 共 {len(board)} 天 | 最新 {doc['last_updated']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
